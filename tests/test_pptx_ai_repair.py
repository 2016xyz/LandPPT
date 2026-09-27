import asyncio
import threading
from io import BytesIO

import fitz
import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from test_template_package_pptx_import import fake_render, fixture_pptx

from landppt.api import template_package_import_api as api
from landppt.services.template_package import pptx_import, pptx_repair
from landppt.services.template_package.builtin import load_builtin_package
from landppt.services.template_package.service import validate_package


@pytest.mark.parametrize("text,size", [("秘", 250), ("正文" * 1001, 24)])
def test_out_of_contract_source_is_retained_for_ai(text, size):
    prs = Presentation(BytesIO(fixture_pptx(1)))
    paragraph = prs.slides[0].shapes[1].text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(size)
    with fitz.open() as doc:
        candidates, warnings = pptx_import._candidates(
            prs, prs.slides[0], doc.new_page(width=960, height=540)
        )
    candidate = next(c for c in candidates if c["text"] == text)
    assert candidate["repair_issues"]
    assert 18 <= candidate["font_size"] <= 160
    assert warnings


def setup_import(monkeypatch):
    monkeypatch.setattr(pptx_import, "render_pdf", fake_render)
    raw = fixture_pptx(1)
    page = pptx_import.analyze_pptx(raw)["slides"][0]
    options = pptx_import.ImportOptions(
        name="Repair test",
        pages=[
            {
                "slide": 1,
                "bindings": {c["id"]: c["suggested"] for c in page["candidates"]},
            }
        ],
    )
    original = pptx_import._component
    saved = {}

    def undersized(*args):
        component = original(*args)
        saved["component"] = component
        root = etree.fromstring(component.svg.encode())
        node = next(n for n in root.iter() if n.get("id") == component.slots[0].node_id)
        saved["patch"] = {
            "node_id": node.get("id"),
            "x": float(node.get("x")),
            "y": float(node.get("y")) - float(node.get("font-size")),
            "width": float(node.get("data-box-w")),
            "height": float(node.get("data-box-h")),
            "font_size": float(node.get("font-size")),
        }
        node.set("data-box-w", "12")
        return component.model_copy(
            update={"svg": etree.tostring(root, encoding="unicode")}
        )

    monkeypatch.setattr(pptx_import, "_component", undersized)
    return raw, options, saved


def test_failed_capacity_is_ai_repaired_and_fully_validated(monkeypatch):
    raw, options, saved = setup_import(monkeypatch)
    calls, events = [], []

    def repair(component, error):
        calls.append(error)
        return pptx_repair.apply_layout_repair(component, {"patches": [saved["patch"]]})

    package, report = pptx_import.import_pptx(
        raw, options, lambda *e: events.append(e), repair
    )
    assert len(calls) == 1 and "只能容纳" in calls[0]
    assert not report["failures"]
    assert any("AI 修复" in w for w in report["warnings"])
    assert any(stage == "repair" for stage, _ in events)
    c = package.components[0]
    assert c.examples[0] == saved["component"].examples[0]
    assert c.slots == saved["component"].slots
    assert c.reference_asset == saved["component"].reference_asset
    validate_package(package)


def test_invalid_patch_gets_second_attempt_and_cannot_change_background(monkeypatch):
    raw, options, saved = setup_import(monkeypatch)
    errors = []

    def repair(component, error):
        errors.append(error)
        patch = dict(saved["patch"])
        if len(errors) == 1:
            patch["node_id"] = "background"
        return pptx_repair.apply_layout_repair(component, {"patches": [patch]})

    package, report = pptx_import.import_pptx(raw, options, repair=repair)
    assert len(errors) == 2 and "不能修改底图" in errors[1]
    assert not report["failures"]
    validate_package(package)


def test_unrepaired_page_never_passes_validation(monkeypatch):
    raw, options, _ = setup_import(monkeypatch)
    calls = []

    def repair(component, error):
        calls.append(error)
        return component

    with pytest.raises(ValueError, match="没有版式通过校验"):
        pptx_import.import_pptx(raw, options, repair=repair)
    assert len(calls) == 2


@pytest.mark.parametrize(
    "field,value", [("font_size", 5), ("width", float("nan")), ("x", 1280)]
)
def test_geometry_patch_rejects_unreadable_nonfinite_and_outside(field, value):
    patch = dict(node_id="title", x=10, y=10, width=600, height=100, font_size=32)
    patch[field] = value
    with pytest.raises(ValueError):
        pptx_repair.GeometryPatch.model_validate(patch)


def test_image_patch_ignores_inapplicable_font_without_changing_bindings():
    component = next(
        c
        for c in load_builtin_package().components
        if any(s.kind == "image" for s in c.slots)
    )
    slot = next(s for s in component.slots if s.kind == "image")
    repaired = pptx_repair.apply_layout_repair(
        component,
        {
            "patches": [
                {
                    "node_id": slot.node_id,
                    "x": 100,
                    "y": 200,
                    "width": 300,
                    "height": 300,
                    "font_size": 24,
                }
            ]
        },
    )
    root = etree.fromstring(repaired.svg.encode())
    node = next(n for n in root.iter() if n.get("id") == slot.node_id)
    assert node.get("y") == "200.0" and node.get("font-size") is None
    assert repaired.slots == component.slots
    assert repaired.examples == component.examples
    with pytest.raises(ValueError, match="不能把内容图片扩大"):
        pptx_repair.apply_layout_repair(
            component,
            {
                "patches": [
                    {
                        "node_id": slot.node_id,
                        "x": 0,
                        "y": 0,
                        "width": 1280,
                        "height": 720,
                    }
                ]
            },
        )


@pytest.mark.asyncio
async def test_import_bridge_uses_current_service_and_cancels_model(monkeypatch):
    started, cancelled = asyncio.Event(), asyncio.Event()
    thread_done = threading.Event()
    service = object()

    async def model(component, error, actual_service):
        assert actual_service is service
        started.set()
        try:
            await asyncio.Future()
        finally:
            cancelled.set()

    def importer(data, settings, progress, repair):
        try:
            repair("component", "overflow")
        finally:
            thread_done.set()

    monkeypatch.setattr(api, "repair_layout", model)
    monkeypatch.setattr(api, "import_pptx", importer)
    task = asyncio.create_task(
        api.import_with_ai_repair(b"", None, service, lambda *_: None)
    )
    await asyncio.wait_for(started.wait(), 3)
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    await asyncio.wait_for(cancelled.wait(), 3)
    assert await asyncio.to_thread(thread_done.wait, 3)
