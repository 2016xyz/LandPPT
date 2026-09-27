import asyncio
import json
from io import BytesIO
from unittest.mock import AsyncMock

import fitz
import httpx
import pytest
from fastapi import FastAPI
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from test_template_package_pptx_import import fake_render, fixture_pptx

from landppt.services.template_package import pptx_analysis, pptx_import


def analysis_page():
    return {
        "slide": 1,
        "candidates": [
            {"id": "1:0", "kind": "text", "text": "页面标题"},
            {"id": "2:0", "kind": "text", "text": "要点标题"},
            {
                "id": "3:0",
                "kind": "text",
                "text": "原主题的小字号正文",
                "font_size": 18,
            },
            {"id": "4:0", "kind": "text", "text": "ANTIMATTER"},
        ],
    }


def plan():
    return {
        "family": "points",
        "assignments": [
            {"id": "1:0", "role": "title", "reason": "页面标题"},
            {"id": "2:0", "role": "heading", "group": "a", "reason": "要点标题"},
            {"id": "3:0", "role": "body", "group": "a", "reason": "属于上方要点的说明"},
            {"id": "4:0", "role": "remove", "reason": "原主题标签"},
        ],
    }


def test_ai_mapping_keeps_small_body_and_groups_it_with_heading():
    slide = analysis_page()
    pptx_analysis.apply_analysis(slide, plan())
    assert slide["candidates"][2]["suggested"] == "body"
    assert slide["candidates"][1]["group"] == slide["candidates"][2]["group"]
    assert slide["candidates"][3]["suggested"] == "remove"


@pytest.mark.parametrize("defect", ["missing", "duplicate", "unknown"])
def test_ai_mapping_rejects_incomplete_or_ambiguous_output(defect):
    value = plan()
    if defect == "missing":
        value["assignments"].pop()
    elif defect == "duplicate":
        value["assignments"].append(value["assignments"][0])
    else:
        value["assignments"][0]["id"] = "invented"
    with pytest.raises(ValueError):
        pptx_analysis.apply_analysis(analysis_page(), value)


def _box_page():
    page = analysis_page()
    page["candidates"].insert(3, {"id": "3:1", "kind": "text", "text": "第二段正文"})
    for index, candidate in enumerate(page["candidates"]):
        candidate["box"] = [100, 100 + index * 60, 400, 40]
    return page


def test_several_bodies_in_one_point_become_separate_points():
    # Real case: a bulleted box under one heading, all tagged with one group.
    value = plan()
    value["assignments"].insert(
        3, {"id": "3:1", "role": "body", "group": "a", "reason": "同一要点的第二段"}
    )
    slide = _box_page()
    pptx_analysis.apply_analysis(slide, value)
    by_id = {c["id"]: c for c in slide["candidates"]}
    assert by_id["2:0"]["suggested"] == "heading"
    assert by_id["3:0"]["group"] == by_id["2:0"]["group"]
    assert by_id["3:1"]["suggested"] == "body"
    assert by_id["3:1"]["group"] not in {"a", ""}


def test_heading_without_body_becomes_a_standalone_point():
    value = plan()
    value["assignments"][2]["group"] = "other"
    slide = analysis_page()
    pptx_analysis.apply_analysis(slide, value)
    assert slide["candidates"][1]["suggested"] == "body"
    assert slide["candidates"][1]["group"] == "a"


def test_normalization_keeps_every_object_and_one_role_per_point():
    items = [
        ("h", "heading", "g", [0, 100, 10, 10]),
        ("h2", "heading", "g", [0, 300, 10, 10]),
        ("b0", "body", "g", [0, 50, 10, 10]),
        ("b1", "body", "g", [0, 140, 10, 10]),
        ("b2", "body", "g", [0, 200, 10, 10]),
    ]
    result = pptx_analysis.normalize_point_groups(items)
    assert set(result) == {i[0] for i in items}
    assert result["h"] == ("heading", "g") and result["b1"] == ("body", "g")
    seen = {}
    for role, group in result.values():
        assert role not in seen.setdefault(group, set())
        seen[group].add(role)
    assert all("body" in roles for roles in seen.values())


def test_long_group_names_get_distinct_suffixes_without_looping():
    name = "g" * 80
    items = [(f"b{i}", "body", name, [0, i * 50, 200, 30]) for i in range(4)]
    result = pptx_analysis.normalize_point_groups(items)
    groups = [group for role, group in result.values()]
    assert len(set(groups)) == 4
    assert all(len(group) <= 80 for group in groups)


def test_numeric_group_labels_match_string_labels():
    result = plan()
    result["assignments"][1]["group"] = 1
    result["assignments"][2]["group"] = "1"
    result["assignments"][0]["group"] = None
    slide = analysis_page()
    pptx_analysis.apply_analysis(slide, result)
    assert slide["candidates"][1]["group"] == slide["candidates"][2]["group"] == "1"


def test_missing_object_diagnostics_name_the_ids():
    result = plan()
    result["assignments"].pop()
    with pytest.raises(ValueError, match="4:0"):
        pptx_analysis.apply_analysis(analysis_page(), result)


def test_decorative_ordinals_do_not_consume_point_capacity():
    slide = {
        "slide": 1,
        "candidates": [
            {"id": "title", "kind": "text", "text": "页面标题"},
            {"id": "number", "kind": "text", "text": "01"},
            *[
                {"id": f"p{i}", "kind": "text", "text": f"要点说明{i}"}
                for i in range(20)
            ],
        ],
    }
    result = {
        "family": "points",
        "assignments": [
            {"id": "title", "role": "title", "reason": "标题"},
            {"id": "number", "role": "body", "group": "number", "reason": "编号"},
            *[
                {"id": f"p{i}", "role": "body", "group": f"g{i}", "reason": "要点"}
                for i in range(20)
            ],
        ],
    }
    pptx_analysis.apply_analysis(slide, result)
    assert slide["candidates"][1]["suggested"] == "fixed"
    assert sum(c["suggested"] == "body" for c in slide["candidates"]) == 20


@pytest.mark.asyncio
async def test_dense_page_keeps_ai_result_for_review_and_continues(monkeypatch):
    dense = {
        "slide": 8,
        "candidates": [
            {"id": "title", "kind": "text", "text": "密集页标题"},
            *[
                {"id": f"p{i}", "kind": "text", "text": f"要点说明{i}"}
                for i in range(21)
            ],
        ],
    }
    result = {
        "family": "points",
        "assignments": [
            {"id": "title", "role": "title", "reason": "标题"},
            *[
                {"id": f"p{i}", "role": "body", "group": f"g{i}", "reason": "要点"}
                for i in range(21)
            ],
        ],
    }
    completion = AsyncMock(side_effect=[(result, {}), (plan(), {})])
    monkeypatch.setattr(pptx_analysis.ContentService, "json_completion", completion)
    analysis = await pptx_analysis.classify_analysis(
        {"slides": [dense, analysis_page()]}, None, lambda *_: None
    )
    assert completion.await_count == 2  # density is not a failed model call
    assert analysis["slides"][0]["review_required"] is True
    assert "21" in analysis["slides"][0]["warnings"][0]
    assert sum(c["suggested"] == "body" for c in dense["candidates"]) == 21
    assert analysis["slides"][1]["analysis_method"] == "ai"


@pytest.mark.asyncio
async def test_transient_model_failure_retries_and_keeps_page_diagnostics(
    monkeypatch, caplog
):
    completion = AsyncMock(side_effect=[TimeoutError("gateway timeout"), (plan(), {})])
    monkeypatch.setattr(pptx_analysis.ContentService, "json_completion", completion)
    result = await pptx_analysis.classify_analysis(
        {"slides": [analysis_page()]}, None, lambda *_: None
    )
    assert result["analysis_method"] == "ai"
    assert "page 1" in caplog.text and "TimeoutError" in caplog.text


@pytest.mark.asyncio
async def test_analysis_cancellation_does_not_retry_model(monkeypatch):
    completion = AsyncMock(side_effect=asyncio.CancelledError)
    monkeypatch.setattr(pptx_analysis.ContentService, "json_completion", completion)
    with pytest.raises(asyncio.CancelledError):
        await pptx_analysis.classify_analysis(
            {"slides": [analysis_page()]}, None, lambda *_: None
        )
    assert completion.await_count == 1


@pytest.mark.asyncio
async def test_ai_analysis_retries_invalid_json_then_reports_page(monkeypatch):
    completion = AsyncMock(side_effect=[ValueError("invalid JSON"), (plan(), {})])
    monkeypatch.setattr(pptx_analysis.ContentService, "json_completion", completion)
    events = []
    result = await pptx_analysis.classify_analysis(
        {"slides": [analysis_page()]}, None, lambda *args: events.append(args)
    )
    assert result["analysis_method"] == "ai"
    assert len(events) == 2
    assert completion.call_args.kwargs["role"] == "template_generation"


def test_tiny_body_is_removed_from_background_and_replaced(monkeypatch):
    prs = Presentation(BytesIO(fixture_pptx(1)))
    prs.slides[0].shapes[1].text_frame.paragraphs[0].font.size = Pt(8)
    prs.slides[0].shapes[1].top = Inches(2.5)
    out = BytesIO()
    prs.save(out)
    rendered_text = []

    def capture(data):
        pdf, fonts = fake_render(data)
        with fitz.open(stream=pdf, filetype="pdf") as doc:
            rendered_text.append(doc[0].get_text())
        return pdf, fonts

    monkeypatch.setattr(pptx_import, "render_pdf", capture)
    slide = pptx_import.analyze_pptx(out.getvalue())["slides"][0]
    body = next(c for c in slide["candidates"] if c["text"] == "Body content 1")
    assert body["suggested"] == "body" and body["source_size"] < 12
    assert body["font_size"] == 18
    package, report = pptx_import.import_pptx(
        out.getvalue(),
        pptx_import.ImportOptions(
            name="Small body",
            pages=[
                {
                    "slide": 1,
                    "bindings": {c["id"]: c["suggested"] for c in slide["candidates"]},
                }
            ],
        ),
    )
    assert not report["failures"]
    assert "Body content 1" not in rendered_text[-1]
    assert package.components[0].examples[0].blocks[0].body == "Body content 1"


def test_unrecognized_text_cannot_be_baked_into_background():
    with fitz.open() as doc:
        page = doc.new_page()
        page.insert_text((30, 30), "Old topic copy")
        with pytest.raises(ValueError, match="底图仍含"):
            pptx_import.audit_background_text(page, ["Brand"])
        pptx_import.audit_background_text(page, ["Old topic copy"])


def test_missing_pdf_style_uses_authored_run_style_and_keeps_candidate():
    prs = Presentation(BytesIO(fixture_pptx(1)))
    paragraph = prs.slides[0].shapes[1].text_frame.paragraphs[0]
    run = paragraph.runs[0]
    run.font.size = Pt(15)
    run.font.name = "Arial"
    run.font.bold = True
    run.font.color.rgb = RGBColor(10, 20, 30)
    paragraph.alignment = PP_ALIGN.RIGHT
    with fitz.open() as doc:
        page = doc.new_page(width=960, height=540)
        candidates, warnings = pptx_import._candidates(prs, prs.slides[0], page)
    body = next(c for c in candidates if c["text"] == "Body content 1")
    assert body["style_source"] == "pptx-default"
    assert body["font_size"] == pytest.approx(20)
    assert body["font"] == "Arial" and body["bold"]
    assert body["color"] == "#0a141e" and body["align"] == "right"
    assert any("继续解析" in warning for warning in warnings)


def test_missing_inherited_style_and_theme_rgb_have_safe_defaults():
    prs = Presentation(BytesIO(fixture_pptx(1)))
    frame = prs.slides[0].shapes[1].text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = "First paragraph"
    paragraph.font.size = None
    paragraph.font.name = None
    paragraph.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
    second = frame.add_paragraph()
    second.text = "Second paragraph"
    with fitz.open() as doc:
        candidates, _ = pptx_import._candidates(
            prs, prs.slides[0], doc.new_page(width=960, height=540)
        )
    first = next(c for c in candidates if c["text"] == paragraph.text)
    second = next(c for c in candidates if c["text"] == second.text)
    assert first["font_size"] == pytest.approx(18)
    assert first["font"] == "Noto Sans CJK SC"
    assert first["color"] == "#182f36" and first["align"] == "left"
    assert second["box"][1] > first["box"][1]


def test_style_fallback_continues_import_and_clears_source_body(monkeypatch):
    inputs = []

    def missing_body_style(data):
        prs = Presentation(BytesIO(data))
        inputs.append([s.text for s in prs.slides[0].shapes if s.has_text_frame])
        # Simulate PDF extraction missing this textbox, while PPTX retains it.
        prs.slides[0].shapes[1].text_frame.clear()
        out = BytesIO()
        prs.save(out)
        return fake_render(out.getvalue())

    monkeypatch.setattr(pptx_import, "render_pdf", missing_body_style)
    prs = Presentation(BytesIO(fixture_pptx(1)))
    prs.slides[0].shapes[1].top = Inches(2.5)
    out = BytesIO()
    prs.save(out)
    raw = out.getvalue()
    slide = pptx_import.analyze_pptx(raw)["slides"][0]
    package, report = pptx_import.import_pptx(
        raw,
        pptx_import.ImportOptions(
            name="Fallback",
            pages=[
                {
                    "slide": 1,
                    "bindings": {c["id"]: c["suggested"] for c in slide["candidates"]},
                }
            ],
        ),
    )
    assert not report["failures"]
    assert "Body content 1" not in inputs[-1]
    assert package.components[0].examples[0].blocks[0].body == "Body content 1"


@pytest.mark.asyncio
async def test_progress_stream_preserves_order_and_propagates_errors():
    from landppt.api.template_package_import_api import stream_operation

    app = FastAPI()

    @app.get("/stream")
    def endpoint(fail: bool = False):
        async def operation(progress):
            progress("render", "正在渲染")
            if fail:
                raise ValueError("第 3 页识别失败")
            progress("ai", "正在识别")
            return {"saved": True}

        return stream_operation(operation)

    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.get("/stream")
        events = [json.loads(line) for line in response.text.splitlines()]
        assert [e["type"] for e in events] == ["progress", "progress", "complete"]
        assert events[-1]["result"]["saved"]
        response = await client.get("/stream?fail=true")
        events = [json.loads(line) for line in response.text.splitlines()]
        assert events[-1] == {"type": "error", "message": "第 3 页识别失败"}
