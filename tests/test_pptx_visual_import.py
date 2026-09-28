import asyncio
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from lxml import etree
from pptx import Presentation
from test_pptx_ai_analysis import analysis_page, plan
from test_template_package_pptx_import import fake_render, fixture_pptx

from landppt.ai.base import ImageContent, TextContent
from landppt.api import template_package_import_api as api
from landppt.services.slide.package_generation.content_service import ContentService
from landppt.services.template_package import pptx_analysis, pptx_import, pptx_vision
from landppt.services.template_package.builtin import load_builtin_package
from landppt.services.template_package.service import validate_package


@pytest.mark.asyncio
async def test_image_is_sent_as_multimodal_content_on_selected_role(monkeypatch):
    monkeypatch.setattr(
        "landppt.services.slide.package_generation.content_service.llm_timeout",
        AsyncMock(return_value=2),
    )
    service = SimpleNamespace(
        _chat_completion_for_role=AsyncMock(
            return_value=SimpleNamespace(
                content='{"ok":true}', model="vision", usage={}
            )
        )
    )
    result, _ = await ContentService(service).json_completion(
        "Read image",
        role="template_generation",
        image_urls=["data:image/png;base64,abc"],
    )
    assert result == {"ok": True}
    call = service._chat_completion_for_role.call_args
    assert call.args == ("template_generation",)
    assert call.kwargs["stream_response"] is True
    parts = call.kwargs["messages"][1].content
    assert isinstance(parts[0], TextContent) and parts[0].text == "Read image"
    assert isinstance(parts[1], ImageContent) and parts[1].image_url["url"].endswith(
        ",abc"
    )


@pytest.mark.asyncio
async def test_visual_classification_receives_actual_page_image(monkeypatch):
    page = analysis_page()
    page["preview"] = "data:image/png;base64,abc"
    completion = AsyncMock(return_value=(plan(), {}))
    monkeypatch.setattr(pptx_analysis, "llm_timeout", AsyncMock(return_value=2))
    monkeypatch.setattr(ContentService, "json_completion", completion)
    result = await pptx_analysis.classify_analysis(
        {"slides": [page]}, None, lambda *_: None, vision=True
    )
    assert result["analysis_method"] == "vision"
    assert completion.call_args.kwargs["image_urls"] == [page["preview"]]


def test_blank_visual_groups_do_not_merge_unrelated_points():
    selected = [
        ({"id": str(i), "text": text, "box": [0, i * 50, 200, 30]}, "body")
        for i, text in enumerate(("第一项内容", "第二项内容", "第三项内容"))
    ]
    assert pptx_vision.merge_visual_groups(selected, {"0": "", "1": " "}) == selected


@pytest.mark.parametrize(
    "addition",
    [
        '<image id="screenshot" data-asset="original"/>',
        '<text x="10" y="10">原主题正文</text>',
        '<image id="original" href="https://example.com/image.png"/>',
    ],
)
def test_visual_design_rejects_raster_shortcut_and_unbound_copy(addition):
    component = load_builtin_package().components[0]
    svg = component.svg.replace("</svg>", addition + "</svg>")
    with pytest.raises(ValueError):
        pptx_vision.apply_visual_design(component, {"svg": svg})


def test_structured_visual_plan_generates_valid_editable_svg():
    package = load_builtin_package()
    component = package.components[0]
    nodes = {
        n.get("id"): n
        for n in etree.fromstring(component.svg.encode()).iter()
        if n.get("id")
    }
    rows = []
    for slot in component.slots:
        node = nodes[slot.node_id]
        size = float(node.get("font-size", "18"))
        rows.append(
            {
                "node_id": slot.node_id,
                "x": float(node.get("x")),
                "y": float(node.get("y")) - size,
                "width": float(node.get("data-box-w")),
                "height": max(float(node.get("data-box-h")), size * 1.4),
                "font_size": size,
            }
        )
    plan = {
        "background": "#F4F1EA",
        "gradient_end": "#FFFFFF",
        "slots": rows,
        "decorations": [
            {
                "kind": "ellipse",
                "x": 10,
                "y": 10,
                "width": 20,
                "height": 20,
                "fill": "#B7462D",
            }
        ],
    }
    rebuilt = pptx_vision.draw_visual_plan(component, plan)
    assert rebuilt.slots == component.slots and rebuilt.examples == component.examples
    validate_package(package.model_copy(update={"components": (rebuilt,)}))
    small_fonts = pptx_vision.draw_visual_plan(
        component, {**plan, "slots": [{**row, "font_size": 16} for row in rows]}
    )
    assert all(
        float(node.get("font-size")) >= 18
        for node in etree.fromstring(small_fonts.svg.encode()).iter()
        if node.get("font-size")
    )
    validate_package(package.model_copy(update={"components": (small_fonts,)}))
    with pytest.raises(ValueError, match="恰好覆盖"):
        pptx_vision.draw_visual_plan(component, {**plan, "slots": rows[:-1]})
    narrow = [{**row, "width": 20, "height": 26} for row in rows]
    with pytest.raises(ValueError, match="至少需要高"):
        pptx_vision.draw_visual_plan(component, {**plan, "slots": narrow})
    # A taller box can be inferred from actual copy only in unused space. It
    # must never grow through the following text slot or a surrounding card.
    long_title = "容量验证" * 10
    expanded = component.model_copy(
        update={
            "examples": (
                component.examples[0].model_copy(update={"title": long_title}),
            )
        }
    )
    grow_rows = [{**row} for row in rows]
    title_index = next(
        i for i, slot in enumerate(component.slots) if slot.field == "title"
    )
    title_row = grow_rows[title_index]
    title_row.update(x=20, y=20, width=200, height=26, font_size=18)
    grown = pptx_vision.draw_visual_plan(
        expanded, {**plan, "decorations": [], "slots": grow_rows}
    )
    first = next(
        n
        for n in etree.fromstring(grown.svg.encode()).iter()
        if n.get("id") == title_row["node_id"]
    )
    assert float(first.get("data-box-h")) >= pptx_vision.text_height(
        long_title, 200, 18
    )
    with pytest.raises(ValueError, match="至少需要高"):
        pptx_vision.draw_visual_plan(
            expanded,
            {
                **plan,
                "slots": grow_rows,
                "decorations": [{"x": 10, "y": 10, "width": 220, "height": 50}],
            },
        )
    next(row for row in grow_rows if row is not title_row).update(x=20, y=60, width=200)
    with pytest.raises(ValueError, match="至少需要高"):
        pptx_vision.draw_visual_plan(
            expanded, {**plan, "decorations": [], "slots": grow_rows}
        )


def test_visual_layout_hints_include_complete_copy_and_measured_heights():
    component = load_builtin_package().components[0]
    hints = pptx_vision.layout_hints(component)
    title = next(item for item in hints if item["field"] == "title")
    assert title["complete_text"] == component.examples[0].title
    assert title["font_size"] == 18
    heights = title["minimum_height_by_width"]
    assert heights["200"] >= heights["1120"] >= 22


def without_background(component):
    root = etree.fromstring(component.svg.encode())
    for node in list(root):
        if node.get("data-asset"):
            root.remove(node)
    return etree.tostring(root, encoding="unicode")


def test_visual_grouping_merges_paragraphs_and_preserves_all_source_text():
    page = analysis_page()
    page["candidates"].append({"id": "3:1", "kind": "text", "text": "第二段完整正文"})
    response = plan()
    response["assignments"].append(
        {"id": "3:1", "role": "body", "group": "a", "reason": "同一卡片"}
    )
    pptx_analysis.apply_analysis(page, response, merge_groups=True)
    selected = []
    groups = {}
    for i, c in enumerate(page["candidates"]):
        c["box"] = [100, 100 + i * 30, 400, 30]
        groups[c["id"]] = c.get("group", "")
        if c["suggested"] not in {"fixed", "remove"}:
            selected.append((c, c["suggested"]))
    merged = pptx_vision.merge_visual_groups(selected, groups)
    body = [c for c, role in merged if role == "body"]
    assert len(body) == 1
    assert body[0]["text"] == "原主题的小字号正文\n第二段完整正文"
    assert page["candidates"][2]["text"] == "原主题的小字号正文"  # Source is immutable.


@pytest.mark.parametrize("with_progress", [True, False])
def test_visual_import_renders_once_retries_validation_and_keeps_content(
    monkeypatch, with_progress
):
    renders, calls, events = [], [], []

    def render(raw):
        renders.append(raw)
        return fake_render(raw)

    monkeypatch.setattr(pptx_import, "render_pdf", render)
    raw = fixture_pptx(1)
    page = pptx_import.analyze_pptx(raw)["slides"][0]
    options = pptx_import.ImportOptions(
        name="Visual",
        mode="visual",
        pages=[
            {
                "slide": 1,
                "bindings": {c["id"]: c["suggested"] for c in page["candidates"]},
            }
        ],
    )
    renders.clear()

    def reconstruct(component, error, reference):
        calls.append((component, error, reference))
        svg = without_background(component).replace(
            'viewBox="0 0 1 1"', 'viewBox="0 0 1280 720"'
        )
        if len(calls) == 1:
            svg = svg.replace('viewBox="0 0 1280 720"', 'viewBox="0 0 1 1"')
        return pptx_vision.apply_visual_design(component, {"svg": svg})

    package, report = pptx_import.import_pptx(
        raw,
        options,
        (lambda *args: events.append(args)) if with_progress else None,
        reconstruct=reconstruct,
    )
    assert (
        len(renders) == 1
    )  # Original screenshot only; no raster background generation.
    assert len(calls) == 2 and "viewBox" in calls[1][1]
    assert all(c[2].startswith("data:image/png;base64,") for c in calls)
    assert not report["failures"]
    result = package.components[0]
    assert result.examples[0] == calls[0][0].examples[0]
    assert result.slots == calls[0][0].slots
    assert result.reference_asset in {a.id for a in package.assets}
    assert "data-asset" not in result.svg
    assert any(stage == "vision" for stage, _ in events) == with_progress
    validate_package(package)


def test_nested_group_text_is_editable_and_removed_from_preserved_background(
    monkeypatch,
):
    prs = Presentation(BytesIO(fixture_pptx(1)))
    slide = prs.slides[0]
    body = slide.shapes[1]
    group = slide.shapes.add_group_shape([body])
    slide.shapes.add_group_shape([group])
    out = BytesIO()
    prs.save(out)
    raw = out.getvalue()
    original_pdf = fake_render(fixture_pptx(1))
    snapshots = []

    def render(data):
        parsed = Presentation(BytesIO(data))
        snapshots.append(
            [
                s.text
                for s, _, _ in pptx_import.walk_shapes(parsed.slides[0].shapes)
                if s.has_text_frame
            ]
        )
        return original_pdf if len(snapshots) == 1 else fake_render(data)

    monkeypatch.setattr(pptx_import, "render_pdf", render)
    # IDs stay attached to nested text nodes; lookup must recurse during removal.
    _, parsed, doc, _ = pptx_import._analyze(raw)
    candidates, _ = pptx_import._candidates(parsed, parsed.slides[0], doc[0])
    doc.close()
    assert any(c["text"] == "Body content 1" for c in candidates)
    options = pptx_import.ImportOptions(
        name="Group",
        pages=[
            {
                "slide": 1,
                "bindings": {
                    c["id"]: (
                        "body" if c["text"] == "Body content 1" else c["suggested"]
                    )
                    for c in candidates
                },
            }
        ],
    )
    package, report = pptx_import.import_pptx(raw, options)
    assert not report["failures"]
    assert "Body content 1" not in snapshots[-1]
    assert package.components[0].examples[0].blocks[0].body == "Body content 1"


@pytest.mark.asyncio
async def test_visual_api_bridge_passes_reference_and_current_user_service(monkeypatch):
    service = object()
    callback = AsyncMock(return_value="recreated")
    monkeypatch.setattr(api, "llm_timeout", AsyncMock(return_value=2))
    monkeypatch.setattr(api, "reconstruct_layout", callback)

    def importer(data, settings, progress, repair, reconstruct):
        return reconstruct("page", "capacity", "data:image/png;base64,abc"), {}

    monkeypatch.setattr(api, "import_pptx", importer)
    result = await api.import_with_ai_repair(
        b"", SimpleNamespace(mode="visual"), service, lambda *_: None
    )
    assert result == ("recreated", {})
    callback.assert_awaited_once_with(
        "page", "capacity", "data:image/png;base64,abc", service
    )


@pytest.mark.asyncio
async def test_reconstruction_does_not_suppress_cancellation(monkeypatch):
    async def cancel(*args, **kwargs):
        raise asyncio.CancelledError()

    monkeypatch.setattr(ContentService, "json_completion", cancel)
    with pytest.raises(asyncio.CancelledError):
        await pptx_vision.reconstruct_layout(
            load_builtin_package().components[0], "", "data:image/png;base64,abc", None
        )
