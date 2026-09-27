"""Opt-in real model/LibreOffice verification; saves artifacts, never user projects."""

import asyncio
import json
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from landppt.services.service_instances import get_ppt_service_for_user
from landppt.services.slide.package_generation.content_service import ContentService
from landppt.services.slide.package_generation.renderer import render_page
from landppt.services.template_package.pptx_analysis import classify_analysis
from landppt.services.template_package.pptx_import import (
    ImportOptions,
    analyze_pptx,
    import_pptx,
)


def fixture():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    def text(x, y, width, height, value, size, bold=False):
        shape = slide.shapes.add_textbox(
            Inches(x), Inches(y), Inches(width), Inches(height)
        )
        paragraph = shape.text_frame.paragraphs[0]
        paragraph.text = value
        paragraph.font.name = "Noto Sans CJK SC"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor(24, 47, 54)

    text(0.7, 0.5, 11, 0.8, "反物质探测原理", 32, True)
    for x, heading, body in (
        (0.8, "磁场探测", "利用磁场观察带电粒子的运动方向，识别粒子性质。"),
        (7.0, "辐射探测", "记录粒子产生的辐射信号，分析粒子的质量与能量。"),
    ):
        text(x, 2.2, 4.7, 0.5, heading, 20, True)
        text(x, 2.85, 4.7, 1.5, body, 8)
    text(0.7, 6.8, 4, 0.3, "ANTIMATTER REPORT", 8)
    out = BytesIO()
    prs.save(out)
    return out.getvalue()


async def main(user_id):
    out = Path("artifacts/pptx-ai-import")
    out.mkdir(parents=True, exist_ok=True)
    raw = fixture()
    (out / "source.pptx").write_bytes(raw)

    def progress(stage, message):
        print(stage, message, flush=True)

    analysis = await asyncio.to_thread(analyze_pptx, raw, progress)
    service = get_ppt_service_for_user(user_id)
    analysis = await classify_analysis(analysis, service, progress)
    slide = analysis["slides"][0]
    options = ImportOptions(
        name="AI 导入验证",
        pages=[
            {
                "slide": 1,
                "family": slide["family"],
                "bindings": {c["id"]: c["suggested"] for c in slide["candidates"]},
                "groups": {
                    c["id"]: c["group"] for c in slide["candidates"] if c.get("group")
                },
            }
        ],
    )
    package, report = await asyncio.to_thread(import_pptx, raw, options, progress)
    assert not report["failures"], report
    component = package.components[0]
    outline = {
        "slide_id": "geology",
        "title": "峡谷地质成因",
        "layout": component.id,
        "content_points": ["板块碰撞促使地壳隆升", "河流长期侵蚀塑造峡谷"],
    }
    valid, errors, _ = await ContentService(service).expand(
        {"outline": {"slides": [outline]}}, package, [outline]
    )
    assert not errors, errors
    rendered = render_page(package, component.id, valid["geology"])
    (out / "geology.html").write_text(rendered.html_content, encoding="utf-8")
    (out / "report.json").write_text(
        json.dumps(
            {
                "candidates": [
                    {k: v for k, v in c.items() if k != "preview"}
                    for c in slide["candidates"]
                ],
                "slots": [s.model_dump() for s in component.slots],
                "content": valid["geology"].model_dump(),
                "import_report": report,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(
        "PASS: real AI analysis, LibreOffice backgrounds, new-topic content and render",
        flush=True,
    )


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--user-id",
        type=int,
        required=True,
        help="Uses this user's configured models; incurs model calls",
    )
    args = parser.parse_args()
    asyncio.run(main(args.user_id))
